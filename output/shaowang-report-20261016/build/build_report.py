from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BUILD = Path(__file__).resolve().parent
ROOT = BUILD.parent
SOURCE = ROOT / 'source' / 'astra-pack'
OUT = ROOT / 'deliverables'
OUT.mkdir(exist_ok=True)
DOSSIER = json.loads((BUILD / 'dossier-sections.json').read_text(encoding='utf-8'))
FIGURES = {f['id']: f for f in json.loads((BUILD / 'figures.json').read_text(encoding='utf-8'))}
TABLE_SPECS = {f['id']: f for f in json.loads((BUILD / 'tables.json').read_text(encoding='utf-8'))}
OLD_TABLES = json.loads((BUILD / 'old-tables.json').read_text(encoding='utf-8'))
OLD_BLOCKS = json.loads((BUILD / 'old-blocks.json').read_text(encoding='utf-8'))
CHECKLIST = (SOURCE / 'backfill-checklist.md').read_text(encoding='utf-8')
PH = {}
for line in CHECKLIST.split('## 二、')[0].splitlines():
    if re.match(r'\| \d+ \|', line):
        cells = [s.strip() for s in line.strip('|').split('|')]
        PH[int(cells[0])] = cells[-1]
SHOTS = dict(re.findall(r'(S-\d\d)：([^】]+)', CHECKLIST))
SHOT_FULL = {k: f'【截图待补 {k}：{v}】' for k, v in SHOTS.items()}
SHOT_TITLES = ['态势大屏','情报详情卡','知识图谱视图','问答侧栏与推理步骤','版本自查页面','智能体任务树','插件与 MCP 状态页面','Claude Code 通过 MCP 调用 check_stack','命令行工具在持续集成中的使用']


def md_tables(section: str) -> list[list[list[str]]]:
    tables, current = [], []
    for line in DOSSIER[section].splitlines() + ['']:
        if line.startswith('|'):
            if not re.match(r'^\|[- :|]+$', line):
                current.append([s.strip().replace('`', '') for s in line.strip('|').split('|')])
        elif current:
            tables.append(current)
            current = []
    return tables


TABLES = {f'T-{i:02d}': OLD_TABLES[i] for i in range(1, 16)}
TABLES['T-02'] = md_tables('D4')[0]
TABLES['T-04'] = md_tables('D6')[0]
TABLES['T-06'] = md_tables('D10')[0]
TABLES['T-07'] = md_tables('D18')[0]
TABLES['T-16'] = md_tables('D18')[1]
TABLES['T-17'] = [
    ['字段情形','处理规则','保留信息'],
    ['在野利用','只有 KEV 设置；取并集，不撤销','其他来源仅作传闻提示'],
    ['一级来源撤回','任一一级来源撤回即 withdrawn','撤回来源与时间'],
    ['一级版本区间互斥','取并集，标 conflict=true','原区间与自查分歧提示'],
    ['别名指向两个 CVE','不自动合并，生成 DUPLICATE_OF 候选','两个编号及复核依据'],
]
TABLES['T-01'] = [
    ['对比维度','通用漏洞资料','情报平台路线','普通 RAG 路线','哨望设计'],
    ['AI 组件覆盖','按来源收录规则','通过连接器组织','取决于语料','九类来源，承接无 CVE 公告'],
    ['多源分歧','保留各源描述','依配置与流程处理','依检索与生成实现','逐字段裁决并保留冲突'],
    ['关系依据','来源自带引用','结构化关系组织','文本证据召回','关系理由与置信度必填'],
    ['版本判断','OSV 提供区间结构','需结合具体实现','需额外精确约束','依赖文件与生态区间求解'],
    ['问答引用','提供事实材料','需结合问答实现','生成与检索结合','显式多跳，返回前校验'],
    ['开发接入','API 等入口','依部署与扩展','依实现而定','MCP、OpenAPI、CLI、插件'],
]
TABLES['T-15'] = [
    ['对比项','通用漏洞资料','情报平台路线','普通 RAG 路线','哨望设计'],
    ['无编号公告','依来源覆盖','依连接器适配','依语料范围','P0 五分钟轮询与内部标识'],
    ['时效观测','来源时间字段','依部署观测','依实现观测','五段时间戳，P50/P95'],
    ['冲突裁决','保留各源主张','依研判流程','依生成判断','字段规则与冲突留痕'],
    ['关联解释','来源引用','图关系组织','文本关联','边带理由、方法与置信度'],
    ['多跳问答','提供事实材料','需结合检索生成','需多步检索机制','显式计划与引用校验'],
    ['故障恢复','外部来源条件','依部署与扩展','依运行平台','退避、死信、熔断、降级'],
    ['编码智能体','API 等方式','依集成方式','依实现方式','七工具设计与插件包'],
]
TABLES['T-03'][8][1] = '全量比对'
TABLES['T-03'][10][-1] = '≤24 h（日更类，单列）'
TABLES['T-03'][-1][2] = '1 周；ATLAS 每日检查'
TABLES['T-08'][1][1] = '计划采用 docker compose 单机部署'
TABLES['T-08'][3][1] = '计划境内直连；记录实际运行网络'
TABLES['T-08'][4][1] = '建议国产大模型 API；bge-m3 与 bge-reranker；实际版本待登记'
for row in TABLES['T-07'][1:]:
    if row[0].startswith('NVD'):
        row[2] = '所测时点支持直连方案；持续运行待验证'
    if row[0] == 'arXiv export API':
        row[2] = '改以 RSS 为主；境内 RSS 连通性待测'

PENDING = '【实录待补】'
APP_TABLES = {
 'A-1': ('境内入口实测初稿', [
    ['来源或入口','类别','结果','日期与环境','处置'],
    ['NVD／GitHub／OSV／KEV／EPSS／huntr／ATLAS','主要来源','HTTP 200','2026-10-14，境内','当时支持直连'],
    ['CNNVD 首页；TC260','国内来源／标准','200','2026-10-14，境内','数据接口／标准文件待补'],
    ['CNVD 三入口','国内来源','521','2026-10-14，境内','人工或半自动导入'],
    ['arXiv export API','论文','502','2026-10-14，境内','RSS 主通道，境内待测'],
    ['Quake 查询接口','测绘','调用成功','2026-10-14，境内','另做语法校准'],
    ['Hunter 首页','测绘','200','2026-10-14，境内','业务 API 待测'],
    ['FOFA','测绘','820031','2026-10-14，境内','积分不足'],
    ['ZoomEye','测绘','鉴权未核实','2026-10-14，境内','暂不启用'],
 ]),
 'A-2': ('境外补充观测与证据边界', [
    ['入口','结果','采集环境与时间','证据状态'],
    ['arXiv RSS','200','境外，2026-10-15 15:57','dossier 记录'],
    ['Semantic Scholar bulk','429','境外，2026-10-15 15:57','无凭据限流记录'],
    ['avid-db 仓库','200','境外，2026-10-15 15:57','dossier 记录'],
    ['Hunter 首页','403','境外，2026-10-15 15:57','不能替代境内 API 结果'],
    ['境内原始日志与哈希',PENDING,'需境内复跑','原终端输出未存档'],
 ]),
 'B-1': ('Quake 查询快照记录', md_tables('D18')[1]),
 'B-2': ('组件语法校准台账', [
    ['组件','平台','校准结果','时间与证据'],
    ['Ollama','Quake','44／50，88%，verified','2026-10-15 19:55；哈希待补'],
    *[[n,'待测平台',PENDING,'未采集实录'] for n in ['vLLM','Dify','ComfyUI','Gradio','Open WebUI','NVIDIA Triton','Langflow','Flowise','Ray Dashboard']],
 ]),
 'C-1': ('真实与演示样本阶段追踪模板', [
    ['记录项','真实第三方样本','演示注入样本'],
    ['身份与来源',PENDING,PENDING],
    ['source_published_at',PENDING,PENDING],
    ['first_seen_at',PENDING,PENDING],
    ['normalized_at',PENDING,PENDING],
    ['enriched_at',PENDING,PENDING],
    ['queryable_at',PENDING,PENDING],
    ['Critic／冲突／关系证据',PENDING,PENDING],
    ['统计归属','漏洞类首次发现按协议计入','DEMO 排除'],
 ]),
 'D-1': ('客户端调用实录登记模板', [
    ['客户端与路径','工具加载','只读策略','长任务续取','记录文件'],
    *[[n,PENDING,PENDING,PENDING,PENDING] for n in ['Claude Code','Codex 远程','Codex stdio','Cursor 远程','Gemini CLI stdio','CLI／CI']],
 ]),
 'E-1': ('性能运行环境登记模板', [
    ['项目','实际运行记录'],
    *[[n,PENDING] for n in ['CPU／内存／磁盘','系统与服务版本','模型／嵌入／重排版本','网络环境','数据规模','代码与配置版本','开始结束时间／观测天数']],
 ]),
 'E-2': ('来源时延原始记录索引', [
    ['来源','样本与时间精度','CSV／复算脚本','结果与哈希'],
    *[[n,PENDING,PENDING,PENDING] for n in ['P0 厂商公告','GHSA','NVD','OSV','huntr','日更来源（单列）','DEMO（排除）']],
 ]),
 'FAPP-1': ('已有失败与方案调整记录', [
    ['事件','已有观察','调整','仍缺什么'],
    ['CNVD','三入口 521','降低主次，人工导入','境内原始日志'],
    ['arXiv','境内 export 502','RSS 主通道','境内 RSS 实测'],
    ['Quake','宽计数约 82 万','收窄后抽样 88%','原始证据附件与哈希'],
    ['FOFA','820031','积分足够后启用','计数实测'],
    ['Hunter','境外首页 403','等待境内 API 测试','业务接口证据'],
    ['ZoomEye','鉴权未核实','暂不启用','官方鉴权核对'],
 ]),
 'FAPP-2': ('故障注入运行登记模板', [
    ['场景','实际转移与恢复','数据完整性','日志／时间／哈希'],
    *[[n,PENDING,PENDING,PENDING] for n in ['429 限流','连续 503','三周期 5xx','400 参数错误','暴露面维度超时','终止 workers']],
 ]),
 'G-1': ('数据集标注登记模板', [
    ['数据集或类别','设计内容','实际标注记录'],
    ['合并／版本／ATLAS','200 条金标，含困难样本',PENDING],
    ['双人标注与仲裁','κ 目标 ≥0.8',PENDING],
    ['问答四题型','约 100 题，含答案与证据',PENDING],
    ['数据集版本与哈希','冻结样本与修订记录',PENDING],
 ]),
 'G-2': ('评测运行与错例登记模板', [
    ['记录项','向量 RAG','BM25＋向量','三路＋Critic'],
    *[[n,PENDING,PENDING,PENDING] for n in ['代码版本／模型版本','配置／提示／Schema','数据集／结果文件','运行时间／哈希','错误样本与归因']],
 ]),
}

CODE = {}
current_code = None
for block in OLD_BLOCKS:
    if block['type'] != 'paragraph':
        continue
    if block['style'] == 'CodeCap':
        current_code = f'C-{len(CODE) + 1:02d}'
        CODE[current_code] = {'title': block['text'], 'lines': []}
    elif block['style'] == 'CodeText' and current_code:
        CODE[current_code]['lines'].append(block['text'])
    elif current_code:
        current_code = None

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(2.54)
section.left_margin, section.right_margin = Cm(2.8), Cm(2.6)
section.header_distance = section.footer_distance = Cm(1.2)


def font_style(style, size: float, bold: bool = False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for key, value in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','宋体')]:
        fonts.set(qn(f'w:{key}'), value)
    for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        fonts.attrib.pop(qn(f'w:{key}'), None)


normal = doc.styles['Normal']
font_style(normal, 12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.first_line_indent = Pt(24)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.widow_control = True
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for level, size in [(1,16),(2,14),(3,12)]:
    st = doc.styles[f'Heading {level}']
    font_style(st, size, True)
    st.paragraph_format.first_line_indent = Pt(0)
    st.paragraph_format.space_before = Pt(12 if level > 1 else 0)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.page_break_before = level == 1

for name, size, bold in [('FrontHeading',16,True),('FigureCaption',10.5,False),('TableCaption',10.5,False),('CodeCaption',10.5,False),('CodeText',9,False),('Reference',10.5,False),('SmallNote',10.5,False)]:
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = normal
    font_style(st, size, bold)
    pf = st.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.line_spacing = 1 if name in ['CodeText','Reference'] else 1.15
    pf.space_after = Pt(4 if name != 'CodeText' else 0)
    if name.endswith('Caption') or name == 'FrontHeading':
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.keep_with_next = name != 'FigureCaption'
    if name == 'FrontHeading':
        pf.space_after = Pt(12)
        pf.page_break_before = True
    if name == 'Reference':
        pf.left_indent = Pt(20)
        pf.first_line_indent = Pt(-20)
for name in ['Title','Subtitle']:
    font_style(doc.styles[name], 24 if name == 'Title' else 16, name == 'Title')

for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
for level in range(1,4):
    name = f'TOC {level}'
    st = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    font_style(st, 11)
    st.paragraph_format.line_spacing = 1.15
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.first_line_indent = Pt(0)

settings = doc.settings.element
update = OxmlElement('w:updateFields')
update.set(qn('w:val'), 'true')
settings.append(update)


def field(paragraph, instruction: str, result: str = ''):
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    begin.set(qn('w:dirty'),'true')
    code = OxmlElement('w:instrText'); code.set(qn('xml:space'),'preserve'); code.text = f' {instruction} '
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate')
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    for elem in [begin, code, sep]:
        paragraph.add_run()._r.append(elem)
    paragraph.add_run(result)
    paragraph.add_run()._r.append(end)


header = section.header.paragraphs[0]
header.text = '哨望 ShaoWang 技术报告'
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
header.paragraph_format.first_line_indent = Pt(0)
for run in header.runs:
    run.font.size = Pt(9)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent = Pt(0)
footer.add_run('第 '); field(footer, 'PAGE','1'); footer.add_run(' 页 共 '); field(footer,'NUMPAGES','1'); footer.add_run(' 页')

mapping = {'figures': [], 'tables': [], 'codes': [], 'placeholders': [], 'references': [], 'sections': [], 'backfill': []}
reference_order = []
figure_count, table_count = defaultdict(int), defaultdict(int)
chapter = '0'
location = '封面'
paragraph_number = 0


def write_runs(paragraph, text: str, references: bool = True):
    global paragraph_number
    text = re.sub(r'\{\{PH:(\d+)\}\}', lambda match: PH[int(match.group(1))], text)
    parts = re.split(r'(【(?:待测|截图待补|图待补|实录待补)[^】]*】|\*\*[^*]+\*\*|\[\d+\])', text)
    for piece in parts:
        if not piece:
            continue
        if references and re.fullmatch(r'\[\d+\]', piece):
            old_id = int(piece[1:-1])
            if old_id not in reference_order:
                reference_order.append(old_id)
            run = paragraph.add_run(f'[{reference_order.index(old_id)+1}]')
            run.font.superscript = True
            run.font.size = Pt(9)
        else:
            bold = piece.startswith('**')
            run = paragraph.add_run(piece[2:-2] if bold else piece)
            if bold or piece.startswith('【'):
                run.bold = True
            if piece.startswith('【'):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                mapping['placeholders'].append({'location':location,'text':piece,'paragraph':paragraph_number})
    paragraph_number += 1


def body(text, style=None):
    p = doc.add_paragraph(style=style)
    write_runs(p, text)
    return p


def caption(kind: str, identifier: str, title: str, count: int, chapter_id: str):
    style = 'FigureCaption' if kind == '图' else 'TableCaption'
    p = doc.add_paragraph(style=style)
    separator = '-' if chapter_id.isalpha() else '.'
    p.add_run(f'{kind} {chapter_id}{separator}')
    field(p, f'SEQ {"Fig" if kind == "图" else "Tab"}{chapter_id} \\* ARABIC', str(count))
    p.add_run(f' {title}')
    return f'{kind} {chapter_id}{separator}{count}'


def add_table(identifier, title, rows):
    global location
    section_location = location
    table_count[chapter] += 1
    number = caption('表',identifier,title,table_count[chapter],chapter)
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    ncol = len(rows[0])
    widths = {2:[4.4,11.2],3:[4.0,5.8,5.8],4:[2.7,4.1,4.1,4.7],5:[2.4,3.0,2.8,3.0,4.4]}.get(ncol,[15.6/ncol]*ncol)
    if identifier == 'T-02': widths = [2.7,3.2,3.4,6.3]
    if identifier == 'T-06': widths = [2.0,4.4,6.2,3.0]
    if identifier == 'T-09': widths = [2.1,4.3,2.0,7.2]
    if identifier == 'T-10': widths = [3.4,2.2,2.4,2.4,5.2]
    if identifier == 'T-11': widths = [2.3,3.0,2.2,2.6,5.5]
    if identifier == 'T-16' or identifier == 'B-1': widths = [2.8,5.5,3.2,4.1]
    if identifier in ['T-01','T-15']: widths = [2.5,3.0,3.1,2.8,4.2]
    for col, width in zip(table.columns, widths): col.width = Cm(width)
    borders = OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        element=OxmlElement(f'w:{name}'); element.set(qn('w:val'),'nil'); borders.append(element)
    table._tbl.tblPr.append(borders)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        trpr = table.rows[-1]._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trpr.append(cant)
        if row_index == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for ci, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Cm(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = row_index == 0
            location = f'{section_location}／{number}／第{row_index+1}行第{ci+1}列'
            write_runs(p, value.replace('`',''))
            for run in p.runs:
                run.font.size = Pt(10.5)
                if row_index == 0: run.bold = True
            tcpr=cell._tc.get_or_add_tcPr()
            cell_borders=OxmlElement('w:tcBorders')
            for side, weight in [('top',12 if row_index == 0 else None),('bottom',12 if row_index == len(rows)-1 else 6 if row_index == 0 else None)]:
                if weight:
                    edge=OxmlElement(f'w:{side}');edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),str(weight));edge.set(qn('w:color'),'000000');cell_borders.append(edge)
            tcpr.append(cell_borders)
    location = section_location
    mapping['tables'].append({'id':identifier,'number':number,'title':title,'location':location})
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(identifier):
    figure_count[chapter] += 1
    if identifier.startswith('S-'):
        title = SHOT_TITLES[int(identifier[-2:])-1]
        path = SOURCE/'build'/'fig'/f'{identifier}.png'
        width = 15.5
    else:
        title = FIGURES[identifier]['title']
        path = ROOT/'fig'/f'{identifier}.png'
        width = 10 if FIGURES[identifier]['layout'] == '半栏' else 15.5
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    if identifier == 'F-28':
        write_runs(p,'【图待补：各来源时延分布箱线图，含 6 h 阈值线】')
    else:
        shape = p.add_run().add_picture(str(path),width=Cm(width))
        shape._inline.docPr.set('descr',f'{identifier} {title}；设计图' if identifier.startswith('F-') else f'{identifier} 截图占位，非真实界面')
    number = caption('图',identifier,title,figure_count[chapter],chapter)
    mapping['figures'].append({'id':identifier,'number':number,'title':title,'location':location,'width_cm':width})
    if identifier.startswith('S-'):
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
        body(SHOT_FULL[identifier])
        extras = {'S-01':'真实截图还需包含 ATLAS 热力图与仅计数的暴露面趋势。','S-02':'来源须显示 tier 徽章。','S-03':'线型区分规则、LLM 和待审候选。','S-05':'底部保留等价 CLI。'}
        if identifier in extras: body(extras[identifier], 'SmallNote')


def add_code(identifier):
    entry = CODE[identifier]
    body(entry['title'], 'CodeCaption')
    for index, line in enumerate(entry['lines']):
        p = doc.add_paragraph(style='CodeText')
        p.paragraph_format.keep_with_next = index < len(entry['lines'])-1
        run = p.add_run(line)
        run.font.name='Consolas';run.font.size=Pt(9)
        shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'F2F2F2');p._p.get_or_add_pPr().append(shd)
    mapping['codes'].append({'id':identifier,'title':entry['title'],'location':location,'lines':len(entry['lines'])})


def process_markdown(text):
    global chapter, location
    text = re.sub(r'(?m)^(#{1,3} .+|\{\{[^{}]+\}\})$', r'\n\1\n', text)
    for block in re.split(r'\n\s*\n', text.strip()):
        block = block.strip()
        if not block: continue
        if block.startswith('#'):
            match=re.match(r'^(#{1,3}) (.+)',block)
            level=len(match.group(1)); title=match.group(2)
            location=title
            if level == 1:
                if title.startswith('第'):
                    chapter = str('一二三四五六'.index(title[1])+1)
                elif title.startswith('附录 '): chapter=title.split()[1]
            style=f'Heading {level}'
            if title in ['摘要','Abstract']: style='FrontHeading'
            p=doc.add_paragraph(title,style=style)
            if title.startswith('附录 ') and chapter != 'A':
                p.paragraph_format.page_break_before = False
                p.paragraph_format.space_before = Pt(18)
            mapping['sections'].append({'title':title,'chapter':chapter})
        elif block == '{{TOC}}':
            doc.add_paragraph('目录',style='FrontHeading')
            field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u','目录待 Word 更新')
        elif block == '{{LISTS}}':
            doc.add_paragraph('图目录',style='FrontHeading')
            field(doc.add_paragraph(),'TOC \\h \\z \\t "FigureCaption,1"','图目录待 Word 更新')
            doc.add_paragraph('表目录',style='FrontHeading')
            field(doc.add_paragraph(),'TOC \\h \\z \\t "TableCaption,1"','表目录待 Word 更新')
        elif re.fullmatch(r'\{\{[^{}]+\}\}',block):
            identifier=block[2:-2]
            if identifier.startswith('F-') or identifier.startswith('S-'):add_figure(identifier)
            elif identifier.startswith('T-'):add_table(identifier,TABLE_SPECS[identifier]['title'],TABLES[identifier])
            elif identifier.startswith('C-') and len(identifier)==4:add_code(identifier)
            elif identifier in APP_TABLES:add_table(identifier,*APP_TABLES[identifier])
            else:raise ValueError(identifier)
        else:
            body(block.replace('\n',' '))


cover=doc.add_paragraph('中国电子杯',style='Subtitle');cover.alignment=WD_ALIGN_PARAGRAPH.CENTER
cover.paragraph_format.space_before=Pt(70)
p=doc.add_paragraph('作品报告',style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(45)
for text in ['赛题九','智能体驱动的 AI 安全知识情报系统','ShaoWang','智能体驱动的 AI 安全知识情报系统','多源实时监测、可解释关联推理与开放插件生态']:
    p=doc.add_paragraph(text,style='Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
p=doc.add_paragraph('提交日期：　　　　年　　月　　日');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(70)

for name in ['00','01','02','03','04','05','06']:
    process_markdown((BUILD/'parts'/f'{name}.md').read_text(encoding='utf-8-sig'))

location='参考文献'
doc.add_paragraph('参考文献',style='Heading 1')
refs={int(m.group(1)):m.group(2) for m in re.finditer(r'^(\d+)\. (.+)$',DOSSIER['D21'],re.M)}
for index, original_id in enumerate(reference_order,1):
    p=doc.add_paragraph(style='Reference');p.add_run(f'[{index}] {refs[original_id]}')
    mapping['references'].append({'number':index,'dossier_number':original_id,'text':refs[original_id]})
process_markdown((BUILD/'parts'/'appendices.md').read_text(encoding='utf-8-sig'))

doc.core_properties.title='ShaoWang：智能体驱动的 AI 安全知识情报系统'
doc.core_properties.subject='赛题九技术报告'
doc.core_properties.author=''
doc.core_properties.last_modified_by=''
doc.core_properties.keywords='AI 安全；威胁情报；多智能体；知识图谱；GraphRAG；模型上下文协议'
path=OUT/'ShaoWang-技术报告.docx'
doc.save(path)
(BUILD/'report-map.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2),encoding='utf-8')
(BUILD/'table-content.json').write_text(json.dumps(TABLES,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(path),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'figures':len(mapping['figures']),'references':len(reference_order),'placeholders':len(mapping['placeholders']),'codes':len(CODE)},ensure_ascii=False))
