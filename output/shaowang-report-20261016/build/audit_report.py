from pathlib import Path
from collections import Counter
from zipfile import ZipFile
import json
import re
from docx import Document
from lxml import etree
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
BUILD=ROOT/'build'
doc=Document(ROOT/'deliverables'/'ShaoWang-技术报告.docx')
paragraphs=[]
for el in doc.element.body:
    if el.tag.endswith('}p'):
        from docx.text.paragraph import Paragraph
        p=Paragraph(el,doc)
        if not p.style.name.startswith('toc'):
            paragraphs.append(p.text)
    if el.tag.endswith('}tbl'):
        from docx.table import Table
        paragraphs.extend(c.text for r in Table(el,doc).rows for c in r.cells)
text='\n'.join(paragraphs)
mapping=json.loads((BUILD/'report-map.json').read_text(encoding='utf-8'))
checklist=(ROOT/'source/astra-pack/backfill-checklist.md').read_text(encoding='utf-8')
items=[]
for line in checklist.split('## 二、')[0].splitlines():
    if re.match(r'\| \d+ \|',line):
        cells=[s.strip() for s in line.strip('|').split('|')]
        items.append((int(cells[0]),cells[-1]))
expected=Counter(value for num,value in items if num not in [9,75])
actual=Counter(x['text'] for x in mapping['placeholders'])
missing={key:count-actual[key] for key,count in expected.items() if actual[key]<count}
for shot in range(1,10):
    match=re.search(r'【截图待补 S-'+f'{shot:02d}'+r'：[^】]+】',checklist)
    if match and actual[match.group(0)] != 1:
        missing[f'S-{shot:02d}']=1
ban=['随着','日益普及','越来越受到关注','具有重要意义','理论意义和应用价值','意义重大','填补空白','国内领先','业界首创','首个','革命性','显著提升','极大地提高','大幅降低','有效解决','全面提升','性能优异','效果良好','赋能','助力','打造','一站式','全方位','多维度','深度融合','智慧化','值得注意的是','不难发现','众所周知','总而言之','综上所述','由此可见','咱们','搞定','重要内容预览']
ban_hits={word:text.count(word) for word in ban if word in text}
legacy=['TAGGED','RELATED_PAPER','AISI-','AISV-','三类实时通道','8 个 MCP 工具']
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
codeold=[]
old=json.loads((BUILD/'old-blocks.json').read_text(encoding='utf-8'))
codeold=[x['text'] for x in old if x.get('style')=='CodeText']
codenew=[p.text for p in doc.paragraphs if p.style.name=='CodeText']
fields=doc.element.xpath('//w:instrText')
figs={item['id'] for item in mapping['figures'] if item['id'].startswith('F-')}
tables={item['id'] for item in mapping['tables'] if item['id'].startswith('T-')}
placeholders=Counter(re.match(r'【([^：\s】]+)',p['text']).group(1) for p in mapping['placeholders'])
files={}
for part in (BUILD/'parts').glob('*.md'):
    string=part.read_text(encoding='utf-8-sig')
    files[part.name]={'characters':len(string),'cjk':len(re.findall(r'[\u4e00-\u9fff]',string))}
audit={
 'backfill_missing':missing,'backfill_required_pending_count':sum(expected.values()),
 'placeholder_counts':dict(placeholders),'banned_phrases':ban_hits,
 'legacy_hits':[word for word in legacy if word in text],
 'ipv4_hits':re.findall(r'(?<!\d)(?:\d{1,3}\.){3}\d{1,3}(?!\d)',text),
 'credential_assignments':re.findall(r'(?i)\b(?:api_?key|token|password)\s*=\s*[\"\'][^\"\']+[\"\']',text),
 'token_keyword_hits':len(re.findall('token',text,re.I)),
 'unresolved_builder_markers':re.findall(r'\{\{.*?\}\}',text),
 'code_identical':codeold==codenew,'code_line_count':len(codenew),
 'figure_ids_complete':figs=={f'F-{i:02d}' for i in range(1,31)},
 'table_ids_complete':tables=={f'T-{i:02d}' for i in range(1,18)},
 'reference_count':len(mapping['references']),
 'reference_ids':sorted(x['dossier_number'] for x in mapping['references']),
 'toc_fields':[f.text for f in fields if 'TOC ' in f.text],
 'seq_field_count':sum('SEQ ' in f.text for f in fields),
 'part_lengths':files,
}
(BUILD/'report-audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
(BUILD/'final-report-extracted.txt').write_text(text,encoding='utf-8')
print(json.dumps(audit,ensure_ascii=False,indent=2))
